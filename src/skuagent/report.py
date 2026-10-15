"""Weekly Excel workbook: the format leadership and owning teams actually open."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY, TEAL = "0B1F3A", "0E8C8C"

DETAIL_COLS = ["sku", "vendor", "category", "material", "weight_lb", "ship_mode", "private_label", "orders", "gross",
               "contribution", "cont_pct", "cogs_pct", "freight_pct", "freight_vs_estimate", "refund_pct", "return_freight",
               "dmg_rate", "roas", "primary_driver", "root_cause", "action_id", "action", "owner", "jira_project",
               "fixability", "confidence", "margin_gap", "evidence"]


def _style(ws, widths=None):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col[0].value, 14)


def write_report(path, diag, action_summary, validation, tickets, insights) -> Path:
    path = Path(path)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"Weekly insights": insights.split("\n")}).to_excel(xw, sheet_name="Weekly Insights", index=False)
        action_summary.to_excel(xw, sheet_name="Action Summary", index=False)
        diag[[c for c in DETAIL_COLS if c in diag]].sort_values("margin_gap", ascending=False).to_excel(xw, sheet_name="SKU Detail", index=False)
        if validation is not None:
            validation.to_excel(xw, sheet_name="LLM Validation", index=False)
        if tickets is not None and len(tickets):
            tickets.to_excel(xw, sheet_name="Jira Tickets", index=False)
        wb = xw.book
        _style(wb["Weekly Insights"], {"Weekly insights": 120})
        _style(wb["Action Summary"], {"action": 48, "owner": 18})
        _style(wb["SKU Detail"], {"root_cause": 46, "action": 44, "evidence": 120})
        if "LLM Validation" in wb.sheetnames:
            ws = wb["LLM Validation"]; _style(ws, {"judge_reason": 60})
            for row in ws.iter_rows(min_row=2):
                agree = row[6].value  # action_agree
                row[6].fill = PatternFill("solid", fgColor="D1FAE5" if agree else "FEE2E2")
        if "Jira Tickets" in wb.sheetnames:
            _style(wb["Jira Tickets"], {"summary": 80})
    return path
